"""Bounded data-only CSV/XLSX parsing and formula-safe spreadsheet exports."""
import csv
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
import re
import unicodedata
import zipfile

from defusedxml import ElementTree as SafeET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import pandas as pd

from src.security import (MAX_FILE_BYTES, MAX_ROWS, MAX_COLUMNS, MAX_CELLS,
                          MAX_CELL_CHARS, SecurityError)

MIMES = {
    '.csv': {'text/csv', 'text/plain', 'application/csv', 'application/vnd.ms-excel'},
    '.xlsx': {'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'},
}
MAX_UNPACKED = 80 * 1024 * 1024


def checked_suffix(filename: str, mime: str | None = None) -> str:
    """Client names are validated but never used as a server-side path."""
    if (not isinstance(filename, str) or not filename or len(filename) > 180
            or '/' in filename or '\\' in filename or ':' in filename
            or '..' in filename or any(ord(c) < 32 for c in filename)):
        raise SecurityError('upload_name', 'Недопустимое имя файла.')
    suffix = Path(filename).suffix.lower()
    if suffix not in MIMES:
        raise SecurityError('upload_type', 'Разрешены только CSV и XLSX.')
    # MIME is an additional check, never the source of trust. Some browsers send
    # octet-stream; those files still undergo full content validation below.
    if mime and mime.split(';')[0].lower() not in MIMES[suffix] | {'application/octet-stream'}:
        raise SecurityError('upload_mime', 'Тип содержимого не соответствует CSV/XLSX.')
    return suffix


def validate_size(size: int) -> None:
    if size <= 0 or size > MAX_FILE_BYTES:
        raise SecurityError('upload_size', 'Файл пуст или превышает 20 МБ.', 413)


def _cell(value) -> str:
    text = '' if value is None else str(value)
    if len(text) > MAX_CELL_CHARS or any(ord(c) < 32 and c not in '\t\r\n' for c in text):
        raise SecurityError('cell_size', 'Ячейка слишком длинная или содержит бинарные данные.')
    return text


def _frame(rows) -> pd.DataFrame:
    header = None
    values = []
    for row in rows:
        if not row or all(value is None or value == '' for value in row):
            continue
        if len(row) > MAX_COLUMNS:
            raise SecurityError('columns_limit', f'Допускается до {MAX_COLUMNS} колонок.')
        cells = [_cell(value) for value in row]
        if header is None:
            header = [v.strip() for v in cells]
            if any(not v for v in header) or len(set(header)) != len(header):
                raise SecurityError('headers', 'Заголовки должны быть заполнены и уникальны.')
            continue
        if len(cells) != len(header):
            raise SecurityError('row_shape', 'Количество значений не совпадает с заголовком.')
        if len(values) >= MAX_ROWS or (len(values) + 1) * len(header) > MAX_CELLS:
            raise SecurityError('rows_limit', 'Слишком много строк или ячеек.', 413)
        values.append(cells)
    if header is None or not values:
        raise SecurityError('empty_table', 'Нужны заголовок и хотя бы одна строка данных.')
    return pd.DataFrame(values, columns=header).replace('', pd.NA)


def _csv(data: bytes) -> pd.DataFrame:
    if data.startswith((b'PK', b'MZ', b'\x7fELF')) or b'\0' in data:
        raise SecurityError('csv_content', 'CSV должен содержать текстовую таблицу.')
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:
            text = data.decode('cp1251')  # Common 1C exports.
        except UnicodeDecodeError as exc:
            raise SecurityError('csv_encoding', 'Используйте CSV в UTF-8 или Windows-1251.') from exc
    if text.lstrip().lower().startswith(('<!doctype', '<html', '<?xml', '<script')):
        raise SecurityError('csv_content', 'Ожидается CSV, а не HTML/XML.')
    try:
        delimiter = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|').delimiter
    except csv.Error:
        delimiter = ','  # A one-column reference file is valid.
    return _frame(csv.reader(StringIO(text, newline=''), delimiter=delimiter, strict=True))


def _inspect_archive(data: bytes) -> None:
    """Inspect ZIP members/XML before openpyxl can allocate expanded sheets."""
    with zipfile.ZipFile(BytesIO(data)) as archive:
        infos = archive.infolist()
        names = [entry.filename for entry in infos]
        if len(infos) > 128 or len({n.casefold() for n in names}) != len(names):
            raise SecurityError('xlsx_archive', 'Недопустимая структура XLSX.')
        if not {'[Content_Types].xml', 'xl/workbook.xml'} <= set(names):
            raise SecurityError('xlsx_content', 'Файл не является книгой XLSX.')
        if sum(entry.file_size for entry in infos) > MAX_UNPACKED:
            raise SecurityError('xlsx_expansion', 'Распакованная книга превышает лимит.', 413)
        sheets = 0
        for info in infos:
            name = info.filename
            if (PurePosixPath(name).is_absolute() or '..' in PurePosixPath(name).parts
                    or '\\' in name or info.flag_bits & 1 or info.is_dir()
                    or info.file_size > 40 * 1024 * 1024
                    or info.file_size / max(info.compress_size, 1) > 200):
                raise SecurityError('xlsx_archive', 'Небезопасная структура или сжатие XLSX.')
            lower = name.lower()
            if (not lower.endswith(('.xml', '.rels')) or any(part in lower for part in
                    ('vbaproject', 'activex', 'embeddings', 'externallinks', 'connections', 'querytables', 'customui'))):
                raise SecurityError('xlsx_active', 'Разрешены только таблицы без макросов и встроенных объектов.')
            with archive.open(info) as member:
                for _, element in SafeET.iterparse(member, events=('end',), forbid_dtd=True):
                    tag = element.tag.rsplit('}', 1)[-1]
                    if tag == 'f':
                        raise SecurityError('xlsx_formula', 'Формулы XLSX запрещены. Экспортируйте значения.')
                    if tag == 'Relationship' and element.get('TargetMode', '').lower() == 'external':
                        raise SecurityError('xlsx_external', 'Внешние связи XLSX запрещены.')
                    if tag == 'Override' and any(v in element.get('ContentType', '').lower()
                                                 for v in ('macroenabled', 'vba', 'activex')):
                        raise SecurityError('xlsx_active', 'Активное содержимое XLSX запрещено.')
                    if name == 'xl/workbook.xml' and tag == 'sheet':
                        sheets += 1
                    if lower.startswith('xl/worksheets/'):
                        if tag == 'row' and not 1 <= int(element.get('r', '1')) <= MAX_ROWS + 1:
                            raise SecurityError('xlsx_dimensions', 'Размер листа превышает лимит.', 413)
                        if tag in {'c', 'dimension'} and element.get('r' if tag == 'c' else 'ref'):
                            bounds = range_boundaries(element.get('r' if tag == 'c' else 'ref'))
                            if (any(v is None for v in bounds) or bounds[2] > MAX_COLUMNS
                                    or bounds[3] > MAX_ROWS + 1):
                                raise SecurityError('xlsx_dimensions', 'Размер листа превышает лимит.', 413)
                    if element.text and len(element.text) > MAX_CELL_CHARS:
                        raise SecurityError('cell_size', 'Ячейка слишком длинная.')
                    element.clear()
        if sheets != 1:
            raise SecurityError('xlsx_sheets', 'XLSX должен содержать ровно один лист с данными.')


def parse_table_bytes(data: bytes, suffix: str) -> pd.DataFrame:
    validate_size(len(data))
    try:
        if suffix == '.csv':
            return _csv(data)
        if suffix != '.xlsx':
            raise SecurityError('upload_type', 'Разрешены только CSV и XLSX.')
        _inspect_archive(data)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            return _frame(workbook.worksheets[0].iter_rows(values_only=True))
        finally:
            workbook.close()
    except SecurityError:
        raise
    except Exception as exc:
        raise SecurityError('malformed_table', 'Не удалось прочитать CSV/XLSX. Проверьте формат файла.') from exc


def read_bounded_table(path: str | Path) -> pd.DataFrame:
    """CLI paths are trusted operator input; the web layer never accepts paths."""
    path = Path(path)
    suffix = checked_suffix(path.name)
    if not path.is_file() or path.is_symlink():
        raise SecurityError('input_file', 'Входной файл отсутствует или не является обычным файлом.')
    validate_size(path.stat().st_size)
    with path.open('rb') as handle:
        data = handle.read(MAX_FILE_BYTES + 1)
    return parse_table_bytes(data, suffix)


def spreadsheet_text(value):
    """Neutralize formulas in strings; legitimate numeric cells remain numeric."""
    if not isinstance(value, str):
        return value
    stripped = ''.join(c for c in value if unicodedata.category(c) not in {'Cf', 'Cc'}).lstrip()
    if stripped.startswith(('=', '+', '-', '@')) or value.startswith(('\t', '\r', '\n')):
        return "'" + value
    return value


def spreadsheet_safe(table: pd.DataFrame) -> pd.DataFrame:
    result = table.copy()
    result.columns = [spreadsheet_text(str(col)) for col in result.columns]
    for col in result.select_dtypes(include=['object', 'string']):
        result[col] = result[col].map(spreadsheet_text)
    return result


def to_excel_bytes(table: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    spreadsheet_safe(table).to_excel(buffer, engine='openpyxl', index=False)
    return buffer.getvalue()


def safe_sheet_name(value: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\\x00-\x1f]", '_', value).strip("'")[:31] or 'Поставщик'
    name, counter = base, 1
    while name.casefold() in used:
        suffix = f'_{counter}'
        name = base[:31 - len(suffix)] + suffix
        counter += 1
    used.add(name.casefold())
    return name
