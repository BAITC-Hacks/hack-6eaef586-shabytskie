"""Local defensive regressions, no requests to external infrastructure."""
from io import BytesIO
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

import pandas as pd
from openpyxl import Workbook, load_workbook

from src.config import Config
from src.data_loader import load_data
from src.ordering import to_1c_csv, export_orders
from src.pipeline import prepare
from src.secure_files import (checked_suffix, parse_table_bytes, spreadsheet_safe, to_excel_bytes,
                              safe_sheet_name, MAX_FILE_BYTES)
from src.security import SecurityError, public_error, security_event
from src.web_auth import Principal, principal_from_claims, authorize
from src.web_service import WebService, JobLimiter, apply_approval_edits


def workbook_bytes(rows=None) -> bytes:
    book = Workbook()
    for row in rows or [['sku', 'quantity'], ['001', 10]]:
        book.active.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def actor(role='manager', name='one') -> Principal:
    return Principal(f'issuer|{name}', role, time.time() + 600)


class FileSecurityTests(unittest.TestCase):
    def test_valid_csv_and_xlsx(self):
        for suffix, content in [('.csv', b'sku;quantity\n001;10\n'), ('.xlsx', workbook_bytes())]:
            table = parse_table_bytes(content, suffix)
            self.assertEqual(table.sku.iloc[0], '001')
            self.assertEqual(float(table.quantity.iloc[0]), 10)

    def test_forbidden_extensions_and_mime(self):
        for name in ['sales.xlsm', 'sales.xls', 'payload.exe', 'file.csv.exe']:
            with self.subTest(name=name), self.assertRaises(SecurityError):
                checked_suffix(name)
        with self.assertRaises(SecurityError):
            checked_suffix('sales.csv', 'text/html')

    def test_traversal_names(self):
        for name in ['../../secrets.csv', '/tmp/sales.csv', 'C:\\data\\file.csv', '..\\file.csv', 'a\n.csv']:
            with self.subTest(name=name), self.assertRaises(SecurityError):
                checked_suffix(name)

    def test_binary_or_html_cannot_be_renamed_to_csv(self):
        for content in [b'MZbinary', b'PKnotcsv', b'<html><script>test</script></html>', b'a,b\n\x00,1']:
            with self.subTest(content=content), self.assertRaises(SecurityError):
                parse_table_bytes(content, '.csv')

    def test_oversized_upload(self):
        with self.assertRaises(SecurityError) as error:
            parse_table_bytes(b'x' * (MAX_FILE_BYTES + 1), '.csv')
        self.assertEqual(error.exception.status, 413)

    def test_malformed_xlsx(self):
        with self.assertRaises(SecurityError):
            parse_table_bytes(b'not a zip archive', '.xlsx')

    def test_formulas_in_xlsx_rejected(self):
        with self.assertRaisesRegex(SecurityError, 'Формулы'):
            parse_table_bytes(workbook_bytes([['sku', 'quantity'], ['001', '=1+1']]), '.xlsx')

    def test_external_link_rejected(self):
        buffer = BytesIO(workbook_bytes())
        with zipfile.ZipFile(buffer, 'a') as archive:
            archive.writestr('xl/worksheets/_rels/sheet1.xml.rels',
                             '<Relationships><Relationship TargetMode="External" Target="https://example.invalid"/></Relationships>')
        with self.assertRaisesRegex(SecurityError, 'Внешние'):
            parse_table_bytes(buffer.getvalue(), '.xlsx')

    def test_macros_and_zip_traversal_rejected(self):
        for name in ['xl/vbaProject.bin', '../../outside.xml']:
            buffer = BytesIO(workbook_bytes())
            with zipfile.ZipFile(buffer, 'a') as archive:
                archive.writestr(name, b'not executable')
            with self.subTest(name=name), self.assertRaises(SecurityError):
                parse_table_bytes(buffer.getvalue(), '.xlsx')

    def test_high_compression_archive_rejected(self):
        buffer = BytesIO(workbook_bytes())
        with zipfile.ZipFile(buffer, 'a', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/unused.xml', '<a>' + 'x' * 1_000_000 + '</a>')
        with self.assertRaises(SecurityError):
            parse_table_bytes(buffer.getvalue(), '.xlsx')

    def test_xml_entities_rejected(self):
        buffer = BytesIO(workbook_bytes())
        with zipfile.ZipFile(buffer, 'a') as archive:
            archive.writestr('xl/unused.xml', '<!DOCTYPE a [<!ENTITY x "test">]><a>&x;</a>')
        with self.assertRaises(SecurityError):
            parse_table_bytes(buffer.getvalue(), '.xlsx')

    def test_multiple_sheets_rejected(self):
        book = Workbook()
        book.create_sheet('second')
        output = BytesIO()
        book.save(output)
        with self.assertRaises(SecurityError):
            parse_table_bytes(output.getvalue(), '.xlsx')

    def test_row_column_cell_limits(self):
        with patch('src.secure_files.MAX_ROWS', 1), self.assertRaises(SecurityError):
            parse_table_bytes(b'a,b\n1,2\n3,4', '.csv')
        with patch('src.secure_files.MAX_COLUMNS', 1), self.assertRaises(SecurityError):
            parse_table_bytes(b'a,b\n1,2', '.csv')
        with self.assertRaises(SecurityError):
            parse_table_bytes(('a\n' + 'x' * 2049).encode(), '.csv')

    def test_xlsx_dimension_limit(self):
        book = Workbook()
        book.active.cell(1, 41, 'beyond maximum')
        output = BytesIO()
        book.save(output)
        with self.assertRaises(SecurityError):
            parse_table_bytes(output.getvalue(), '.xlsx')

    def test_csv_formula_injection_is_neutralized(self):
        payloads = ['=1+1', '+SUM(1,1)', '-SUM(1,1)', '@SUM(1,1)', '\t=1+1', '\ufeff=1+1', '  =1+1']
        table = pd.DataFrame({'name': payloads, 'quantity': [-3] * len(payloads)})
        safe = spreadsheet_safe(table)
        self.assertTrue(safe.name.str.startswith("'").all())
        self.assertTrue(safe.quantity.eq(-3).all())
        csv = pd.read_csv(BytesIO(to_1c_csv(table)), sep=';')
        self.assertTrue(csv.name.str.startswith("'").all())
        book = load_workbook(BytesIO(to_excel_bytes(table)), data_only=False)
        self.assertTrue(all(book.active.cell(i, 1).data_type != 'f' for i in range(2, 9)))
        self.assertEqual(book.active.cell(2, 2).value, -3)

    def test_sheet_collision_terminates(self):
        used = {'сводка', 'все позиции'}
        names = [safe_sheet_name('Сводка', used) for _ in range(15)]
        self.assertEqual(len(names), len(set(names)))
        self.assertTrue(all(len(n) <= 31 for n in names))


class ValidationSecurityTests(unittest.TestCase):
    def test_invalid_config(self):
        for values in [{'forecast_days': 99999}, {'service_level_z': float('inf')},
                       {'service_level_z': float('nan')}, {'review_period_days': 0}]:
            with self.subTest(values=values), self.assertRaises(SecurityError):
                Config(**values)

    def test_calendar_amplification_rejected(self):
        frame = pd.DataFrame({'sku': ['A', 'A'], 'date': ['1900-01-01', '2025-01-01'], 'quantity': [1, 1]})
        with self.assertRaises(SecurityError):
            prepare(frame)

    def test_unknown_personal_columns_do_not_reach_audit(self):
        raw = pd.DataFrame({'sku': ['A', 'A'], 'date': ['2025-01-01', '2025-01-02'], 'quantity': [1, 1],
                            'client_id': ['original-id', 'original-id'], 'phone': ['private', 'private']})
        transactions, _ = prepare(raw)
        again, _ = prepare(raw)
        self.assertNotIn('phone', transactions)
        self.assertEqual(transactions.client_id.iloc[0], transactions.client_id.iloc[1])
        self.assertNotEqual(transactions.client_id.iloc[0], again.client_id.iloc[0])

    def test_html_and_sql_like_text_remains_plain_data(self):
        payload = '<script>alert(1)</script>\' OR 1=1 --'
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'sales.csv'
            pd.DataFrame({'sku': [payload], 'quantity': [1], 'date': ['2025-01-01']}).to_csv(path, index=False)
            table = load_data(path)
            self.assertEqual(table.sku.iloc[0], payload)
            self.assertEqual(len(table), 1)
        app = Path('app.py').read_text()
        self.assertNotIn('unsafe_allow_html', app)
        self.assertNotIn('st.html(', app)

    def test_sensitive_error_details_not_exposed_or_logged(self):
        secret = '/private/user/secrets.toml password=do-not-show'
        with self.assertLogs('security', level='INFO') as log:
            message = public_error(RuntimeError(secret))
        self.assertNotIn(secret, message)
        self.assertNotIn(secret, str(log.output))

    def test_event_actor_is_pseudonymous(self):
        with self.assertLogs('security', level='INFO') as log:
            security_event('access_denied', 'private-user@example.invalid')
        self.assertNotIn('private-user', str(log.output))


class AuthorizationTests(unittest.TestCase):
    def claims(self):
        return {'iss': 'https://id.example.invalid', 'sub': 'allowed', 'exp': time.time() + 300}

    def test_missing_authentication_returns_401(self):
        with self.assertRaises(SecurityError) as error:
            authorize(None, 'read')
        self.assertEqual(error.exception.status, 401)

    def test_roles_are_server_allowlisted(self):
        user = principal_from_claims(self.claims(), logged_in=True,
                                     issuer='https://id.example.invalid', roles={'allowed': 'viewer'})
        authorize(user, 'read')
        with self.assertRaises(SecurityError) as error:
            authorize(user, 'approve')
        self.assertEqual(error.exception.status, 403)

    def test_forged_unverified_claims_are_not_authentication(self):
        with self.assertRaises(SecurityError):
            principal_from_claims(self.claims(), logged_in=False,
                                  issuer='https://id.example.invalid', roles={'allowed': 'manager'})

    def test_expired_wrong_issuer_unknown_subject_rejected(self):
        for change in [{'exp': time.time() - 1}, {'iss': 'evil'}, {'sub': 'unknown'}, {'exp': float('nan')}]:
            with self.subTest(change=change), self.assertRaises(SecurityError):
                principal_from_claims(self.claims() | change, logged_in=True,
                                      issuer='https://id.example.invalid', roles={'allowed': 'manager'})

    def test_client_role_claim_is_ignored(self):
        user = principal_from_claims(self.claims() | {'role': 'admin'}, logged_in=True,
                                     issuer='https://id.example.invalid', roles={'allowed': 'viewer'})
        self.assertEqual(user.role, 'viewer')

    def test_public_demo_cannot_upload_or_calculate_real_data(self):
        user = Principal('public-demo', 'demo', float('inf'))
        for action in ['upload', 'calculate', 'approve']:
            with self.subTest(action=action), self.assertRaises(SecurityError):
                authorize(user, action)

    def test_workspace_idor_between_users_and_sessions(self):
        service = WebService()
        try:
            token = service.create(actor(), 'session-one', demo=False)
            for user, session in [(actor(name='two'), 'session-one'), (actor(), 'session-two')]:
                with self.assertRaises(SecurityError) as error:
                    service.resolve(token, user, session)
                self.assertEqual(error.exception.status, 403)
            with self.assertRaises(SecurityError):
                service.resolve('../../outside', actor(), 'session-one')
        finally:
            service.close()

    def test_workspace_retention_and_path_allowlist(self):
        service = WebService()
        try:
            token = service.create(actor(), 'one', demo=False)
            item = service.resolve(token, actor(), 'one')
            root = item.root
            item.ready = True
            with self.assertRaises(SecurityError):
                service.artifact(token, actor(), 'one', '../../secrets')
            item.expires_at = 0
            service.prune()
            self.assertFalse(root.exists())
        finally:
            service.close()

    def test_rate_and_concurrency_limit(self):
        limiter = JobLimiter(per_user=1)
        with limiter.hold('one'):
            with self.assertRaises(SecurityError) as error:
                with limiter.hold('two'):
                    self.fail('parallel job admitted')
            self.assertEqual(error.exception.status, 429)
        with self.assertRaises(SecurityError):
            with limiter.hold('one'):
                self.fail('rate limit bypassed')

    def test_approval_rejects_mass_assignment_and_other_supplier(self):
        orders = pd.DataFrame({'sku': ['A', 'B'], 'supplier': ['S1', 'S2']})
        for edits in [pd.DataFrame({'sku': ['A'], 'approved_qty': [1], 'price': [0]}),
                      pd.DataFrame({'sku': ['B'], 'approved_qty': [1]}),
                      pd.DataFrame({'sku': ['A', 'A'], 'approved_qty': [1, 1]})]:
            with self.subTest(edits=edits), self.assertRaises(SecurityError):
                apply_approval_edits(orders, 'S1', edits)

    def test_worker_timeout_is_safe(self):
        import subprocess
        service = WebService()
        with tempfile.TemporaryDirectory() as directory:
            with patch('src.web_service.subprocess.run', side_effect=subprocess.TimeoutExpired('worker', 120)):
                with self.assertRaises(SecurityError) as error:
                    service._run_worker(Path(directory))
        self.assertEqual(error.exception.code, 'job_timeout')

    def test_cors_csrf_and_traceback_settings(self):
        import tomllib
        config = tomllib.loads(Path('.streamlit/config.toml').read_text())
        self.assertTrue(config['server']['enableCORS'])
        self.assertTrue(config['server']['enableXsrfProtection'])
        self.assertFalse(config['server']['enableStaticServing'])
        self.assertEqual(config['client']['showErrorDetails'], 'none')


if __name__ == '__main__':
    unittest.main()
